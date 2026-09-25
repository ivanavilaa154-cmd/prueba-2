"""Importa exportaciones reales del ERP (CSV o Excel) a una base SQLite local.

Sirve para probar la plataforma con datos reales cuando todavía no hay conexión
directa a la base del ERP. Cada archivo corresponde a una tabla del diccionario
(clientes.csv, ventas.xlsx, ...) y usa sus mismos nombres de columna. Se aceptan
formatos habituales en Uruguay y Argentina: "1.234,56", fechas dd/mm/aaaa,
separador ";" y archivos en Latin-1.

El resultado es una base nueva; los archivos originales no se modifican.
"""
from __future__ import annotations

import csv
import io
import re
import sqlite3
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from .. import config
from .demo import ESQUEMA

RUTA = config.BACKEND / "datos_reales.db"

# Columnas sin las cuales una fila no sirve. El resto puede faltar (queda vacío).
OBLIGATORIAS = {
    "sucursales": ["id"],
    "vendedores": ["id"],
    "proveedores": ["id"],
    "clientes": ["id"],
    "productos": ["id"],
    "stock": ["producto_id", "sucursal_id", "cantidad"],
    "ventas": ["id", "fecha", "cliente_id"],
    "ventas_lineas": ["venta_id", "producto_id", "cantidad", "precio_unitario"],
    "cxc": ["cliente_id", "saldo"],
}
BOOLEANAS = {("ventas", "anulada"), ("productos", "perecedero")}
_VERDADERO = {"1", "si", "sí", "s", "true", "verdadero", "x", "yes"}
_FALSO = {"0", "no", "n", "false", "falso", ""}


class ErrorImportacion(Exception):
    pass


@dataclass
class InformeTabla:
    tabla: str
    archivo: str
    filas: int = 0
    descartadas: int = 0
    columnas_faltantes: list[str] = field(default_factory=list)
    columnas_ignoradas: list[str] = field(default_factory=list)
    problemas: list[str] = field(default_factory=list)


def _esquema() -> dict[str, dict[str, str]]:
    """{tabla: {columna: tipo}} a partir del esquema de la base demo."""
    con = sqlite3.connect(":memory:")
    con.executescript(ESQUEMA)
    tablas = [f[0] for f in con.execute("SELECT name FROM sqlite_master WHERE type = 'table'")]
    esquema = {t: {c[1]: c[2].upper() for c in con.execute(f"PRAGMA table_info({t})")} for t in tablas}
    con.close()
    return esquema


def plantilla(tabla: str) -> str:
    """Encabezado CSV de una tabla, para que el cliente sepa qué columnas exportar."""
    esquema = _esquema()
    if tabla not in esquema:
        raise ErrorImportacion(f"Tabla desconocida: {tabla}")
    return ";".join(esquema[tabla]) + "\n"


def _normalizar(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", texto.strip().lower()).strip("_")


def tabla_de_archivo(nombre: str) -> str | None:
    base = _normalizar(Path(nombre).stem)
    tablas = _esquema()
    if base in tablas:
        return base
    # "Clientes 2026-09" o "export_ventas_lineas" → la tabla más larga que aparezca
    candidatas = [t for t in tablas if re.search(rf"(^|_){t}(_|$)", base)]
    return max(candidatas, key=len) if candidatas else None


def numero(valor, entero: bool = False):
    """Convierte "1.234,56", "1,234.56", "$ 1.234" o "12%" en número."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        return int(valor)
    if isinstance(valor, (int, float)):
        return int(round(valor)) if entero else float(valor)
    texto = str(valor).strip().replace(" ", "").replace(" ", "")
    texto = re.sub(r"^[^\d\-+.,]+", "", texto)  # moneda al principio
    if texto in {"", "-"}:
        return None
    porcentaje = texto.endswith("%")
    texto = texto.rstrip("%")
    if "," in texto and "." in texto:
        decimal = "," if texto.rfind(",") > texto.rfind(".") else "."
        miles = "." if decimal == "," else ","
        texto = texto.replace(miles, "").replace(decimal, ".")
    elif "," in texto:
        texto = texto.replace(",", "") if texto.count(",") > 1 else texto.replace(",", ".")
    elif texto.count(".") > 1 or re.fullmatch(r"-?[1-9]\d{0,2}\.\d{3}", texto):
        texto = texto.replace(".", "")  # "1.234" o "1.234.567": separador de miles
    try:
        resultado = float(texto)
    except ValueError:
        raise ValueError(f"no es un número: {valor!r}") from None
    if porcentaje:
        resultado /= 100
    if entero:
        if resultado != int(resultado):
            raise ValueError(f"se esperaba un número entero: {valor!r}")
        return int(resultado)
    return resultado


def fecha(valor):
    """Devuelve la fecha en formato AAAA-MM-DD (acepta dd/mm/aaaa y fechas de Excel)."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    texto = str(valor).strip().split(" ")[0].split("T")[0]
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"no es una fecha: {valor!r}")


def _convertir(tabla: str, columna: str, tipo: str, valor):
    if isinstance(valor, str):
        valor = valor.strip()
    if (tabla, columna) in BOOLEANAS:
        if isinstance(valor, (int, float)):
            return 1 if valor else 0
        texto = str(valor or "").strip().lower()
        if texto in _VERDADERO:
            return 1
        if texto in _FALSO:
            return 0
        raise ValueError(f"se esperaba sí/no: {valor!r}")
    if columna.startswith("fecha"):
        return fecha(valor)
    if valor is None or valor == "":
        return None
    if tipo == "INTEGER":
        return numero(valor, entero=True)
    if tipo == "REAL":
        return numero(valor)
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))  # códigos numéricos leídos de Excel
    return str(valor)


def _leer_csv(contenido: bytes) -> list[list]:
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:5000]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=";,\t|")
    except csv.Error:
        separador = ";" if muestra.count(";") > muestra.count(",") else ","
        dialecto = type("Separado", (csv.excel,), {"delimiter": separador})
    return [fila for fila in csv.reader(io.StringIO(texto), dialecto) if any(c.strip() for c in fila)]


def _leer_excel(contenido: bytes) -> list[list]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise ErrorImportacion("Para leer Excel instalá openpyxl (pip install -r requirements.txt).") from e
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro.worksheets[0]
    filas = [list(f) for f in hoja.iter_rows(values_only=True)
             if any(c is not None and str(c).strip() != "" for c in f)]
    libro.close()
    return filas


def leer_archivo(nombre: str, contenido: bytes) -> list[list]:
    extension = Path(nombre).suffix.lower()
    if extension in {".xlsx", ".xlsm"}:
        return _leer_excel(contenido)
    if extension in {".csv", ".txt"}:
        return _leer_csv(contenido)
    raise ErrorImportacion(f"{nombre}: formato no soportado (usá .csv o .xlsx).")


def _cargar_tabla(con, tabla: str, columnas: dict[str, str], nombre: str, filas: list[list]) -> InformeTabla:
    informe = InformeTabla(tabla=tabla, archivo=nombre)
    if not filas:
        informe.problemas.append("El archivo está vacío.")
        return informe
    encabezado = [_normalizar(c) if c is not None else "" for c in filas[0]]
    posicion = {c: i for i, c in enumerate(encabezado) if c in columnas}
    informe.columnas_ignoradas = [c for c in encabezado if c and c not in columnas]
    informe.columnas_faltantes = [c for c in columnas if c not in posicion]
    faltan_obligatorias = [c for c in OBLIGATORIAS[tabla] if c not in posicion]
    if faltan_obligatorias:
        raise ErrorImportacion(
            f"{nombre}: faltan columnas obligatorias para {tabla}: {', '.join(faltan_obligatorias)}. "
            f"Columnas esperadas: {', '.join(columnas)}."
        )

    registros, vistos = [], set()
    for numero_fila, fila in enumerate(filas[1:], start=2):
        try:
            registro = {}
            for columna, tipo in columnas.items():
                i = posicion.get(columna)
                valor = fila[i] if i is not None and i < len(fila) else None
                registro[columna] = _convertir(tabla, columna, tipo, valor) if i is not None else None
            vacias = [c for c in OBLIGATORIAS[tabla] if registro[c] is None]
            if vacias:
                raise ValueError("vacío en " + ", ".join(vacias))
            if "id" in OBLIGATORIAS[tabla]:
                if registro["id"] in vistos:
                    raise ValueError(f"id repetido ({registro['id']})")
                vistos.add(registro["id"])
        except ValueError as e:
            informe.descartadas += 1
            if len(informe.problemas) < 10:
                informe.problemas.append(f"Fila {numero_fila}: {e}")
            continue
        registros.append(tuple(registro[c] for c in columnas))

    marcas = ",".join("?" * len(columnas))
    con.executemany(f"INSERT INTO {tabla} ({','.join(columnas)}) VALUES ({marcas})", registros)
    informe.filas = len(registros)
    if informe.descartadas > 10:
        informe.problemas.append(f"... y {informe.descartadas - 10} filas más descartadas.")
    return informe


_REFERENCIAS = [
    ("ventas", "cliente_id", "clientes", "ventas de clientes que no están en clientes"),
    ("cxc", "cliente_id", "clientes", "facturas a cobrar de clientes que no están en clientes"),
    ("ventas_lineas", "venta_id", "ventas", "líneas de ventas que no están en ventas"),
    ("ventas_lineas", "producto_id", "productos", "líneas de productos que no están en productos"),
    ("clientes", "vendedor_id", "vendedores", "clientes con un vendedor que no está en vendedores"),
    ("stock", "producto_id", "productos", "filas de stock de productos que no están en productos"),
]


def _controles(con) -> list[str]:
    avisos = []
    cantidades = {t: con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in OBLIGATORIAS}
    for tabla, columna, destino, texto in _REFERENCIAS:
        if not cantidades[tabla] or not cantidades[destino]:
            continue
        huerfanas = con.execute(
            f"SELECT COUNT(*) FROM {tabla} WHERE {columna} IS NOT NULL "
            f"AND {columna} NOT IN (SELECT id FROM {destino})"
        ).fetchone()[0]
        if huerfanas:
            avisos.append(f"{huerfanas} {texto}.")
    if cantidades["clientes"]:
        sin_vendedor = con.execute("SELECT COUNT(*) FROM clientes WHERE vendedor_id IS NULL").fetchone()[0]
        if sin_vendedor:
            avisos.append(f"{sin_vendedor} clientes sin vendedor asignado: ningún vendedor los va a ver.")
    for tabla, cantidad in cantidades.items():
        if not cantidad:
            avisos.append(f"Sin datos de {tabla}: las preguntas que la necesiten van a responder que falta el dato.")
    return avisos


def importar(archivos: list[tuple[str, bytes]], ruta: Path | None = None) -> dict:
    """Crea la base con los archivos recibidos [(nombre, contenido)] y devuelve el informe.

    Si algo falla no se toca la base anterior.
    """
    ruta = ruta or RUTA
    esquema = _esquema()
    por_tabla: dict[str, tuple[str, bytes]] = {}
    no_reconocidos = []
    for nombre, contenido in archivos:
        tabla = tabla_de_archivo(nombre)
        if tabla is None:
            no_reconocidos.append(nombre)
        elif tabla in por_tabla:
            raise ErrorImportacion(f"Hay dos archivos para {tabla}: {por_tabla[tabla][0]} y {nombre}.")
        else:
            por_tabla[tabla] = (nombre, contenido)
    if no_reconocidos:
        raise ErrorImportacion(
            "No reconozco a qué tabla corresponde: " + ", ".join(no_reconocidos)
            + ". Nombrá cada archivo como la tabla: " + ", ".join(esquema) + "."
        )
    if "clientes" not in por_tabla:
        raise ErrorImportacion("Falta el archivo de clientes (es el mínimo para empezar).")

    temporal = ruta.with_suffix(".tmp")
    temporal.unlink(missing_ok=True)
    con = sqlite3.connect(temporal)
    try:
        con.executescript(ESQUEMA)
        informes = [
            _cargar_tabla(con, tabla, esquema[tabla], nombre, leer_archivo(nombre, contenido))
            for tabla, (nombre, contenido) in por_tabla.items()
        ]
        con.commit()
        avisos = _controles(con)
    except Exception:
        con.close()
        temporal.unlink(missing_ok=True)
        raise
    con.close()
    temporal.replace(ruta)
    return {
        "ruta": str(ruta),
        "tablas": [vars(i) for i in informes],
        "avisos": avisos,
    }
