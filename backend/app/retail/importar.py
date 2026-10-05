"""Importación de archivos (sección 5.2): Excel o CSV de ventas, stock, productos, compras, listas de precios y pedidos de delivery.

Pedidos de delivery (PedidosYa, Rappi): sus APIs solo se habilitan a integradores aprobados, así que entran por el reporte de
pedidos que cada comercio descarga de su panel (una fila por producto del pedido, con la comisión y el envío del pedido).

1. Analizar: se leen los encabezados, se propone qué columna es cada dato (o se usa el mapeo recordado para ese formato de
   archivo) y se guarda todo en staging. Si el mismo archivo ya se importó (misma huella), se avisa.
2. Validar: se aplica el mapeo y se revisa fila por fila (qué falta, qué número no es un número, qué producto no existe).
3. Confirmar: se importa lo válido. Es idempotente: las ventas se identifican por sucursal + número de ticket, así que
   reimportar el mismo archivo (o uno que se superpone) no duplica ventas.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo

from . import catalogo, db

CAMPOS: dict[str, list[tuple[str, str, bool]]] = {
    # tipo: [(campo, etiqueta, obligatorio)]
    "ventas": [("fecha", "Fecha", True), ("hora", "Hora", False), ("sucursal", "Sucursal", True), ("ticket", "Número de ticket", True),
               ("producto", "Producto (código, EAN o nombre)", True), ("cantidad", "Cantidad", True), ("precio", "Precio unitario cobrado", True),
               ("precio_lista", "Precio de lista", False), ("descuento", "Descuento de la línea", False), ("costo", "Costo unitario", False),
               ("canal", "Canal", False), ("cajero", "Cajero", False), ("medio_pago", "Medio de pago", False)],
    "stock": [("producto", "Producto (código, EAN o nombre)", True), ("sucursal", "Sucursal o depósito", True), ("cantidad", "Stock", True),
              ("lote", "Lote", False), ("vencimiento", "Vencimiento", False)],
    "productos": [("codigo", "Código interno", True), ("nombre", "Nombre", True), ("ean", "Código de barras", False), ("marca", "Marca", False),
                  ("categoria", "Categoría", False), ("subcategoria", "Subcategoría", False), ("proveedor", "Proveedor", False),
                  ("costo", "Costo", False), ("precio", "Precio de venta", False), ("bulto", "Unidades por bulto", False), ("unidad", "Unidad (u/kg)", False)],
    "compras": [("fecha", "Fecha", True), ("proveedor", "Proveedor", True), ("documento", "Número de factura o remito", True),
                ("sucursal", "Sucursal que recibe", True), ("producto", "Producto", True), ("cantidad", "Cantidad", True), ("costo", "Costo unitario", True),
                ("lote", "Lote", False), ("vencimiento", "Vencimiento", False)],
    "delivery": [("fecha", "Fecha", True), ("hora", "Hora", False), ("plataforma", "Plataforma (PedidosYa o Rappi)", True),
                 ("pedido", "Número de pedido", True), ("sucursal", "Sucursal (local)", True), ("producto", "Producto (código, EAN o nombre)", True),
                 ("cantidad", "Cantidad", True), ("precio", "Precio unitario", True), ("comision", "Comisión del pedido", False),
                 ("envio", "Envío a cargo del local", False), ("estado", "Estado del pedido", False)],
    "precios": [("proveedor", "Proveedor", True), ("producto", "Código o EAN del producto", False), ("descripcion", "Descripción", False),
                ("costo", "Costo", True), ("vigencia", "Vigente desde", False)],
}
SINONIMOS = {
    "fecha": ["fecha", "dia", "date", "fecha venta", "fecha comprobante"], "hora": ["hora", "time"],
    "sucursal": ["sucursal", "local", "tienda", "deposito", "ubicacion", "punto de venta", "pdv", "store"],
    "ticket": ["ticket", "comprobante", "nro ticket", "numero", "nro", "factura", "orden", "pedido"],
    "producto": ["producto", "codigo", "cod", "ean", "sku", "articulo", "item", "codigo barras", "cod barras", "plu"],
    "cantidad": ["cantidad", "cant", "unidades", "qty", "stock", "existencia", "saldo"], "precio": ["precio", "precio unitario", "p unit", "importe unitario", "pu", "precio venta"],
    "precio_lista": ["precio lista", "lista"], "descuento": ["descuento", "desc", "bonificacion"], "costo": ["costo", "costo unitario", "cost", "precio compra", "neto"],
    "canal": ["canal"], "cajero": ["cajero", "vendedor", "usuario", "operador"], "medio_pago": ["medio pago", "forma pago", "pago", "medio"],
    "lote": ["lote", "partida"], "vencimiento": ["vencimiento", "vence", "fecha vencimiento", "vto"], "codigo": ["codigo", "cod", "codigo interno", "sku", "plu"],
    "nombre": ["nombre", "descripcion", "producto", "articulo", "detalle"], "ean": ["ean", "codigo barras", "cod barras", "barcode", "gtin"],
    "marca": ["marca", "brand"], "categoria": ["categoria", "rubro", "familia"], "subcategoria": ["subcategoria", "subrubro", "subfamilia"],
    "proveedor": ["proveedor", "supplier"], "bulto": ["bulto", "unidades por bulto", "u x bulto", "pack"], "unidad": ["unidad", "um", "unidad medida"],
    "documento": ["documento", "factura", "remito", "comprobante", "numero"], "descripcion": ["descripcion", "detalle", "producto", "articulo"],
    "vigencia": ["vigencia", "vigente desde", "desde", "fecha"],
    "plataforma": ["plataforma", "app", "aplicacion", "canal venta", "marketplace"], "pedido": ["pedido", "orden", "nro pedido", "order id", "id pedido"],
    "comision": ["comision", "comision plataforma", "fee", "cargo servicio"], "envio": ["envio", "costo envio", "delivery fee", "cargo envio"],
    "estado": ["estado", "status", "estado pedido"],
}
DELIVERY = {"pedidosya": ("pedidosya", "PedidosYa"), "pedidos ya": ("pedidosya", "PedidosYa"), "rappi": ("rappi", "Rappi")}
ANULADOS = {"cancelado", "cancelada", "rechazado", "rechazada", "cancelled", "canceled", "rejected"}
MEDIOS = {"efectivo": "efectivo", "contado": "efectivo", "debito": "debito", "tarjeta de debito": "debito", "credito": "credito",
          "tarjeta de credito": "credito", "tarjeta": "credito", "qr": "qr", "mercado pago": "qr", "transferencia": "transferencia",
          "cuenta corriente": "cuenta_corriente", "fiado": "cuenta_corriente"}


def _norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "").lower())
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def leer_archivo(nombre: str, contenido: bytes) -> tuple[list[str], list[list]]:
    """Encabezados y filas de un .xlsx o .csv (separador ; o , detectado; UTF-8 o Latin-1)."""
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hoja = libro.active
        filas = [list(f) for f in hoja.iter_rows(values_only=True)]
    else:
        try:
            texto = contenido.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = contenido.decode("latin-1")
        muestra = texto[:5000]
        separador = ";" if muestra.count(";") >= muestra.count(",") else ","
        filas = [f for f in csv.reader(io.StringIO(texto), delimiter=separador)]
    filas = [f for f in filas if any(c not in (None, "") for c in f)]
    if not filas:
        raise ValueError("El archivo está vacío.")
    encabezados = [str(c).strip() if c is not None else f"Columna {i + 1}" for i, c in enumerate(filas[0])]
    return encabezados, [list(f) + [None] * (len(encabezados) - len(f)) for f in filas[1:]]


def firma(encabezados: list[str]) -> str:
    return "|".join(sorted(_norm(e) for e in encabezados))


def _clave(texto) -> str:
    """Para comparar encabezados: sin tildes, sin signos y sin «de/del/la» («Código de barras» = «codigo barras»)."""
    return " ".join(p for p in _norm(texto).split() if p not in ("de", "del", "la", "el", "x"))


def sugerir_mapeo(tipo: str, encabezados: list[str]) -> dict[str, str | None]:
    usados = set()
    mapeo = {}
    normal = {_clave(e): e for e in encabezados}
    for campo, _, _ in CAMPOS[tipo]:
        elegido = None
        for sin in SINONIMOS.get(campo, [campo]):
            if sin in normal and normal[sin] not in usados:
                elegido = normal[sin]
                break
        if not elegido:
            for n, original in normal.items():
                if original not in usados and any(s in n for s in SINONIMOS.get(campo, [campo]) if len(s) > 3):
                    elegido = original
                    break
        mapeo[campo] = elegido
        if elegido:
            usados.add(elegido)
    return mapeo


def numero(valor) -> Decimal | None:
    """«1.234,56», «1234.56», «$ 1.234», «-3» → Decimal. Vacío → None. Lo que no es número → ValueError."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))
    t = str(valor).strip().replace("$", "").replace(" ", "")
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", t):
        t = t.replace(".", "")
    try:
        return Decimal(t)
    except InvalidOperation:
        raise ValueError(f"«{valor}» no es un número")


def fecha(valor) -> date | None:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)):            # número de serie de Excel
        return date(1899, 12, 30) + timedelta(days=int(valor))
    t = str(valor).strip()[:10]
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%Y"):
        try:
            return datetime.strptime(t, formato).date()
        except ValueError:
            continue
    raise ValueError(f"«{valor}» no es una fecha (usá DD/MM/AAAA)")


def hora(valor) -> time:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return time(12, 0)
    if isinstance(valor, time):
        return valor
    if isinstance(valor, datetime):
        return valor.time()
    if isinstance(valor, float) and valor < 1:
        minutos = round(valor * 24 * 60)
        return time(minutos // 60 % 24, minutos % 60)
    m = re.match(r"(\d{1,2}):(\d{2})", str(valor))
    if not m:
        raise ValueError(f"«{valor}» no es una hora (HH:MM)")
    return time(int(m.group(1)) % 24, int(m.group(2)))


# ------------------------------------------------------------------------------ analizar / validar / confirmar
def analizar(conn, ctx, tipo: str, nombre: str, contenido: bytes) -> dict:
    if tipo not in CAMPOS:
        raise ValueError("Tipo de importación no válido.")
    huella = hashlib.sha256(contenido).hexdigest()
    previo = db.fila(conn, "SELECT id, created_at, filas_ok FROM lotes_importacion WHERE huella=%s AND tipo=%s AND estado='importado'", (huella, tipo))
    encabezados, filas = leer_archivo(nombre, contenido)
    if len(filas) > 200000:
        raise ValueError("El archivo tiene más de 200.000 filas: dividilo en partes (por mes, por ejemplo).")
    f = firma(encabezados)
    recordado = db.fila(conn, "SELECT mapeo FROM mapeos_columnas WHERE tipo=%s AND firma=%s", (tipo, f))
    lote = db.fila(conn, "INSERT INTO lotes_importacion (org_id, tipo, origen, nombre_archivo, huella, estado, filas_total, created_by, encabezados) "
                         "VALUES (%s,%s,'archivo',%s,%s,'pendiente',%s,%s,%s) RETURNING id",
                   (ctx.org_id, tipo, nombre, huella, len(filas), ctx.usuario_id, json.dumps(encabezados, ensure_ascii=False)))
    db.copiar(conn, "staging_filas", ["org_id", "lote_id", "numero", "datos"],
              [(ctx.org_id, lote["id"], i + 2, json.dumps(dict(zip(encabezados, [_serializable(v) for v in fila])), ensure_ascii=False))
               for i, fila in enumerate(filas)])
    return {"lote_id": lote["id"], "encabezados": encabezados, "muestra": [dict(zip(encabezados, [_serializable(v) for v in f])) for f in filas[:8]],
            "filas": len(filas), "campos": [{"campo": c, "etiqueta": e, "obligatorio": o} for c, e, o in CAMPOS[tipo]],
            "mapeo": recordado["mapeo"] if recordado else sugerir_mapeo(tipo, encabezados), "mapeo_recordado": bool(recordado),
            "ya_importado": {"lote": previo["id"], "fecha": previo["created_at"], "filas": previo["filas_ok"]} if previo else None}


def retomar(conn, lote_id: int) -> dict:
    """El mismo resultado de analizar() para un lote que quedó pendiente (por ejemplo, uno que subió el agente con un formato nuevo)."""
    lote = db.fila(conn, "SELECT id, tipo, estado, filas_total, encabezados FROM lotes_importacion WHERE id=%s", (lote_id,))
    if not lote or lote["estado"] != "pendiente" or not lote["encabezados"]:
        raise ValueError("Ese archivo ya no está esperando: subilo de nuevo.")
    encabezados = lote["encabezados"]
    recordado = db.fila(conn, "SELECT mapeo FROM mapeos_columnas WHERE tipo=%s AND firma=%s", (lote["tipo"], firma(encabezados)))
    muestra = [f["datos"] for f in db.filas(conn, "SELECT datos FROM staging_filas WHERE lote_id=%s ORDER BY numero LIMIT 8", (lote_id,))]
    return {"lote_id": lote_id, "tipo": lote["tipo"], "encabezados": encabezados, "muestra": muestra, "filas": lote["filas_total"],
            "campos": [{"campo": c, "etiqueta": e, "obligatorio": o} for c, e, o in CAMPOS[lote["tipo"]]],
            "mapeo": recordado["mapeo"] if recordado else sugerir_mapeo(lote["tipo"], encabezados), "mapeo_recordado": bool(recordado),
            "ya_importado": None}


def _serializable(v):
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return str(v)
    return v


def _filas(conn, lote_id: int) -> list[tuple[int, dict]]:
    return [(f["numero"], f["datos"]) for f in db.filas(conn, "SELECT numero, datos FROM staging_filas WHERE lote_id=%s ORDER BY numero", (lote_id,))]


def _contexto_catalogo(conn) -> dict:
    return {"productos": db.filas(conn, "SELECT id, nombre, marca, codigo_interno, ean FROM productos WHERE activo"),
            "ubicaciones": {_norm(u["nombre"]): u["id"] for u in db.filas(conn, "SELECT id, nombre FROM ubicaciones")} |
                           {_norm(u["codigo_externo"]): u["id"] for u in db.filas(conn, "SELECT id, codigo_externo FROM ubicaciones WHERE codigo_externo IS NOT NULL")},
            "proveedores": {_norm(p["razon_social"]): p["id"] for p in db.filas(conn, "SELECT id, razon_social FROM proveedores")},
            "canales": {c["codigo"]: c["id"] for c in db.filas(conn, "SELECT id, codigo FROM canales")}}


def _producto(conn, cat: dict, valor, origen="archivo", origen_id=None) -> tuple[int | None, float]:
    texto = str(valor or "").strip()
    if not texto:
        return None, 0.0
    memo = cat.setdefault("memo", {})
    clave = (texto, origen, origen_id)
    if clave not in memo:
        memo[clave] = _producto_sin_memo(conn, cat, texto, origen, origen_id)
    return memo[clave]


def _producto_sin_memo(conn, cat: dict, texto: str, origen, origen_id) -> tuple[int | None, float]:
    es_codigo = bool(re.fullmatch(r"[\w\-.]+", texto)) and (texto.isdigit() or not " " in texto)
    e = catalogo.emparejar(conn, ean=texto if texto.isdigit() and len(texto) in (8, 12, 13, 14) else None,
                           codigo=texto if es_codigo else None, texto=None if es_codigo else texto, origen=origen, origen_id=origen_id,
                           productos=cat["productos"])
    return e.producto_id, e.confianza


def validar(conn, ctx, lote_id: int, mapeo: dict[str, str | None]) -> dict:
    lote = db.fila(conn, "SELECT * FROM lotes_importacion WHERE id=%s", (lote_id,))
    if not lote:
        raise ValueError("No existe esa importación.")
    tipo = lote["tipo"]
    faltan = [e for c, e, o in CAMPOS[tipo] if o and not mapeo.get(c)]
    if faltan:
        raise ValueError(f"Falta indicar qué columna es: {', '.join(faltan)}.")
    cat = _contexto_catalogo(conn)
    errores = []
    ok = 0
    sin_producto: dict[str, int] = defaultdict(int)
    with conn.cursor() as cur:
        for numero_fila, datos in _filas(conn, lote_id):
            valor = {c: datos.get(col) if col else None for c, col in mapeo.items()}
            error = _validar_fila(conn, tipo, valor, cat, sin_producto)
            cur.execute("UPDATE staging_filas SET error=%s WHERE lote_id=%s AND numero=%s", (error, lote_id, numero_fila))
            if error:
                if len(errores) < 500:
                    errores.append({"fila": numero_fila, "error": error})
            else:
                ok += 1
        # Se recuerda el mapeo para la próxima vez que llegue un archivo con las mismas columnas.
        primera = db.fila(conn, "SELECT datos FROM staging_filas WHERE lote_id=%s ORDER BY numero LIMIT 1", (lote_id,))
        encabezados = list(primera["datos"].keys()) if primera else []
        cur.execute("INSERT INTO mapeos_columnas (org_id, tipo, firma, mapeo) VALUES (%s,%s,%s,%s) ON CONFLICT (org_id, tipo, firma) "
                    "DO UPDATE SET mapeo=EXCLUDED.mapeo, updated_at=now()", (ctx.org_id, tipo, firma(encabezados), json.dumps(mapeo)))
        cur.execute("UPDATE lotes_importacion SET estado=%s, filas_ok=%s, filas_error=%s, errores=%s, mapeo=%s WHERE id=%s",
                    ("validado" if ok else "con_errores", ok, lote["filas_total"] - ok, json.dumps(errores[:200]), json.dumps(mapeo), lote_id))
    return {"lote_id": lote_id, "filas": lote["filas_total"], "validas": ok, "con_error": lote["filas_total"] - ok, "errores": errores[:200],
            "productos_no_encontrados": sorted(sin_producto.items(), key=lambda x: -x[1])[:30]}


def _validar_fila(conn, tipo: str, v: dict, cat: dict, sin_producto: dict) -> str | None:
    try:
        for c, e, o in CAMPOS[tipo]:
            if o and (v.get(c) is None or str(v.get(c)).strip() == ""):
                return f"Falta {e.lower()}"
        if tipo in importar_distribuidor.CAMPOS:
            return importar_distribuidor.validar_fila(conn, tipo, v, cat, sin_producto)
        if tipo == "delivery" and _plataforma_delivery(v["plataforma"]) is None:
            return f"«{v['plataforma']}» no es PedidosYa ni Rappi"
        if tipo in ("ventas", "compras", "stock", "delivery"):
            if _norm(v["sucursal"]) not in cat["ubicaciones"]:
                return f"No existe la sucursal «{v['sucursal']}» (cargala en Configuración o corregí el nombre)"
        if tipo in ("ventas", "compras", "stock", "delivery"):
            pid, _ = _producto(conn, cat, v["producto"])
            if not pid:
                sin_producto[str(v["producto"])] += 1
                return f"No se encontró el producto «{v['producto']}» (importá primero los productos o emparejalo en Catálogo)"
        if tipo in ("ventas", "compras", "delivery"):
            fecha(v["fecha"])
        if tipo == "delivery":
            hora(v.get("hora"))
            numero(v["cantidad"])
            numero(v["precio"])
            numero(v.get("comision"))
            numero(v.get("envio"))
        if tipo == "ventas":
            hora(v.get("hora"))
            numero(v["cantidad"])
            numero(v["precio"])
            numero(v.get("descuento"))
            numero(v.get("costo"))
        if tipo in ("stock",):
            numero(v["cantidad"])
            fecha(v.get("vencimiento"))
        if tipo == "compras":
            numero(v["cantidad"])
            numero(v["costo"])
            fecha(v.get("vencimiento"))
        if tipo == "productos":
            numero(v.get("costo"))
            numero(v.get("precio"))
            b = numero(v.get("bulto"))
            if b is not None and b <= 0:
                return "Las unidades por bulto tienen que ser mayores a cero"
        if tipo == "precios":
            if numero(v["costo"]) is None:
                return "Falta el costo"
            fecha(v.get("vigencia"))
            if not (v.get("producto") or v.get("descripcion")):
                return "Falta el código o la descripción del producto"
    except ValueError as e:
        return str(e)
    return None


def confirmar(conn, ctx, lote_id: int) -> dict:
    lote = db.fila(conn, "SELECT * FROM lotes_importacion WHERE id=%s", (lote_id,))
    if not lote:
        raise ValueError("No existe esa importación.")
    if lote["estado"] == "importado":
        raise ValueError("Esta importación ya se confirmó.")
    if lote["estado"] != "validado" or not lote["mapeo"]:
        raise ValueError("Primero hay que validar el archivo (y que tenga al menos una fila válida).")
    mapeo = lote["mapeo"]
    filas = [(f["numero"], {c: f["datos"].get(col) if col else None for c, col in mapeo.items()})
             for f in db.filas(conn, "SELECT numero, datos FROM staging_filas WHERE lote_id=%s AND error IS NULL ORDER BY numero", (lote_id,))]
    cat = _contexto_catalogo(conn)
    zona = ZoneInfo(db.fila(conn, "SELECT zona_horaria FROM organizaciones WHERE id = app_org()")["zona_horaria"])
    resultado = {"ventas": _ventas, "stock": _stock, "productos": _productos, "compras": _compras, "precios": _precios,
                 "delivery": _delivery, **importar_distribuidor.MANEJADORES}[lote["tipo"]](conn, ctx, lote, filas, cat, zona)
    with conn.cursor() as cur:
        cur.execute("UPDATE lotes_importacion SET estado='importado', filas_ok=%s, filas_duplicadas=%s, resultado=%s WHERE id=%s",
                    (resultado.get("importadas", 0), resultado.get("duplicadas", 0), json.dumps(resultado, default=str), lote_id))
    return resultado


def _ventas(conn, ctx, lote, filas, cat, zona) -> dict:
    """Idempotente: cada ticket se identifica por sucursal + número; si ya existe (de esta u otra importación), no se vuelve a cargar."""
    tickets: dict = defaultdict(list)
    for _, v in filas:
        uid = cat["ubicaciones"][_norm(v["sucursal"])]
        tickets[(uid, str(v["ticket"]).strip())].append(v)
    importados = duplicados = lineas = 0
    fechas = set()
    with conn.cursor() as cur:
        for (uid, numero_ticket), vs in tickets.items():
            clave = f"{uid}-{numero_ticket}"
            if db.fila(conn, "SELECT 1 FROM tickets WHERE origen='archivo' AND numero_externo=%s", (clave,)):
                duplicados += 1
                continue
            v0 = vs[0]
            f = fecha(v0["fecha"])
            fechas.add(f)
            fh = datetime.combine(f, hora(v0.get("hora")), zona)
            canal = cat["canales"].get(_norm(v0.get("canal") or "fisico").replace(" ", ""), cat["canales"].get("fisico"))
            total = Decimal(0)
            detalle = []
            for v in vs:
                pid, _ = _producto(conn, cat, v["producto"])
                cant = numero(v["cantidad"])
                precio = numero(v["precio"])
                desc = numero(v.get("descuento")) or Decimal(0)
                lista = numero(v.get("precio_lista")) or (precio + desc / cant if cant else precio)
                costo = numero(v.get("costo"))
                if costo is None:
                    c = db.fila(conn, "SELECT costo FROM producto_proveedores WHERE producto_id=%s ORDER BY principal DESC LIMIT 1", (pid,))
                    costo = c["costo"] if c else None
                total += precio * cant
                detalle.append((pid, cant, lista, precio, desc, costo))
            t = db.fila(conn, "INSERT INTO tickets (org_id, ubicacion_id, canal_id, punto_venta, cajero, fecha_hora, total, estado, numero_externo, origen, lote_importacion_id) "
                              "VALUES (%s,%s,%s,NULL,%s,%s,%s,%s,%s,'archivo',%s) ON CONFLICT (org_id, origen, numero_externo) DO NOTHING RETURNING id",
                        (ctx.org_id, uid, canal, v0.get("cajero"), fh, total, "devuelto" if total < 0 else "confirmado", clave, lote["id"]))
            if not t:
                duplicados += 1
                continue
            for pid, cant, lista, precio, desc, costo in detalle:
                cur.execute("INSERT INTO tickets_lineas (org_id, ticket_id, ubicacion_id, fecha, producto_id, cantidad, precio_lista, precio_cobrado, descuento, costo_unitario) "
                            "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)", (ctx.org_id, t["id"], uid, f, pid, cant, lista, precio, desc, costo))
                lineas += 1
            medio = MEDIOS.get(_norm(v0.get("medio_pago") or ""), None)
            if medio:
                cur.execute("INSERT INTO pagos (org_id, ticket_id, ubicacion_id, medio, monto) VALUES (%s,%s,%s,%s,%s)", (ctx.org_id, t["id"], uid, medio, total))
            importados += 1
    return {"importadas": importados, "duplicadas": duplicados, "lineas": lineas, "desde": min(fechas) if fechas else None,
            "hasta": max(fechas) if fechas else None,
            "mensaje": f"{importados} tickets nuevos ({lineas} líneas)" + (f"; {duplicados} ya estaban cargados y no se duplicaron" if duplicados else "")}


def _plataforma_delivery(texto) -> tuple[str, str] | None:
    t = _norm(texto or "").replace("-", " ")
    return next((v for k, v in DELIVERY.items() if k in t or k.replace(" ", "") in t.replace(" ", "")), None)


def _delivery(conn, ctx, lote, filas, cat, zona) -> dict:
    """Idempotente: cada pedido se identifica por plataforma + número. La comisión y el envío son del pedido (se toma el primer
    valor informado, no la suma de las filas). Crea la plataforma y activa el canal delivery la primera vez."""
    pedidos: dict = defaultdict(list)
    for _, v in filas:
        tipo, _nombre = _plataforma_delivery(v["plataforma"])
        pedidos[(tipo, str(v["pedido"]).strip())].append(v)
    plataformas: dict = {}
    importados = duplicados = lineas = anulados = 0
    fechas = set()
    with conn.cursor() as cur:
        canal = db.fila(conn, "SELECT id FROM canales WHERE codigo='delivery'")
        if not canal:
            canal = db.fila(conn, "INSERT INTO canales (org_id, codigo, nombre, created_by) VALUES (%s,'delivery','Delivery',%s) RETURNING id",
                            (ctx.org_id, ctx.usuario_id))
        cur.execute("UPDATE canales SET activo=true WHERE id=%s", (canal["id"],))
        for (tipo, numero_pedido), vs in pedidos.items():
            v0 = vs[0]
            uid = cat["ubicaciones"][_norm(v0["sucursal"])]
            if (tipo, uid) not in plataformas:
                p = db.fila(conn, "SELECT id FROM plataformas WHERE tipo=%s AND ubicacion_despacho_id=%s", (tipo, uid)) or \
                    db.fila(conn, "INSERT INTO plataformas (org_id, tipo, nombre, canal_id, ubicacion_despacho_id, created_by) VALUES (%s,%s,%s,%s,%s,%s) RETURNING id",
                            (ctx.org_id, tipo, _plataforma_delivery(v0["plataforma"])[1], canal["id"], uid, ctx.usuario_id))
                plataformas[(tipo, uid)] = p["id"]
            plataforma_id = plataformas[(tipo, uid)]
            clave = f"{plataforma_id}-{numero_pedido}"
            if db.fila(conn, "SELECT 1 FROM tickets WHERE origen=%s AND numero_externo=%s", (tipo, clave)):
                duplicados += 1
                continue
            f = fecha(v0["fecha"])
            fechas.add(f)
            fh = datetime.combine(f, hora(v0.get("hora")), zona)
            anulado = _norm(v0.get("estado") or "") in ANULADOS
            detalle, total = [], Decimal(0)
            for v in vs:
                pid, _ = _producto(conn, cat, v["producto"])
                cant, precio = numero(v["cantidad"]), numero(v["precio"])
                c = db.fila(conn, "SELECT costo FROM producto_proveedores WHERE producto_id=%s ORDER BY principal DESC LIMIT 1", (pid,))
                detalle.append((pid, cant, precio, c["costo"] if c else None))
                total += cant * precio
            comision = next((numero(v.get("comision")) for v in vs if numero(v.get("comision")) is not None), Decimal(0))
            envio = next((numero(v.get("envio")) for v in vs if numero(v.get("envio")) is not None), Decimal(0))
            t = db.fila(conn, "INSERT INTO tickets (org_id, ubicacion_id, canal_id, plataforma_id, punto_venta, fecha_hora, total, estado, numero_externo, origen, "
                              "lote_importacion_id) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT (org_id, origen, numero_externo) DO NOTHING RETURNING id",
                        (ctx.org_id, uid, canal["id"], plataforma_id, tipo, fh, total, "anulado" if anulado else "confirmado", clave, tipo, lote["id"]))
            if not t:
                duplicados += 1
                continue
            for pid, cant, precio, costo in detalle:
                cur.execute("INSERT INTO tickets_lineas (org_id, ticket_id, ubicacion_id, fecha, producto_id, cantidad, precio_lista, precio_cobrado, descuento, costo_unitario) "
                            "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,0,%s)", (ctx.org_id, t["id"], uid, f, pid, cant, precio, precio, costo))
                lineas += 1
            cur.execute("INSERT INTO pagos (org_id, ticket_id, ubicacion_id, medio, monto, comision_estimada, plazo_acreditacion_dias) VALUES (%s,%s,%s,'plataforma',%s,%s,7)",
                        (ctx.org_id, t["id"], uid, total, comision))
            cur.execute("INSERT INTO costos_canal (org_id, plataforma_id, ticket_id, fecha, comision, envio) VALUES (%s,%s,%s,%s,%s,%s)",
                        (ctx.org_id, plataforma_id, t["id"], f, comision, envio))
            cur.execute("INSERT INTO pedidos_online (ticket_id, org_id, ubicacion_id, plataforma_id, estado, creado_at, despachado_at) VALUES (%s,%s,%s,%s,%s,%s,%s)",
                        (t["id"], ctx.org_id, uid, plataforma_id, "cancelado" if anulado else "entregado", fh, None if anulado else fh))
            importados += 1
            anulados += anulado
    return {"importadas": importados, "duplicadas": duplicados, "lineas": lineas, "anulados": anulados, "desde": min(fechas) if fechas else None,
            "hasta": max(fechas) if fechas else None,
            "mensaje": f"{importados} pedidos de delivery nuevos ({lineas} líneas" + (f", {anulados} cancelados" if anulados else "") + ")"
                       + (f"; {duplicados} ya estaban cargados y no se duplicaron" if duplicados else "")}


def _stock(conn, ctx, lote, filas, cat, zona) -> dict:
    n = 0
    with conn.cursor() as cur:
        for _, v in filas:
            uid = cat["ubicaciones"][_norm(v["sucursal"])]
            pid, _ = _producto(conn, cat, v["producto"])
            cant = numero(v["cantidad"])
            anterior = db.fila(conn, "SELECT cantidad FROM stock_actual WHERE producto_id=%s AND ubicacion_id=%s", (pid, uid))
            dif = cant - (anterior["cantidad"] if anterior else Decimal(0))
            cur.execute("INSERT INTO stock_actual (org_id, producto_id, ubicacion_id, cantidad) VALUES (%s,%s,%s,%s) ON CONFLICT (producto_id, ubicacion_id) "
                        "DO UPDATE SET cantidad=EXCLUDED.cantidad, actualizado_at=now()", (ctx.org_id, pid, uid, cant))
            if dif:
                cur.execute("INSERT INTO movimientos_stock (org_id, producto_id, ubicacion_id, tipo, cantidad, documento_tipo, documento_id, motivo, usuario_id) "
                            "VALUES (%s,%s,%s,%s,%s,'importacion',%s,'Stock importado',%s)",
                            (ctx.org_id, pid, uid, "inicial" if not anterior else "ajuste", dif, lote["id"], ctx.usuario_id))
            if v.get("lote") or v.get("vencimiento"):
                cur.execute("DELETE FROM stock_lotes WHERE producto_id=%s AND ubicacion_id=%s AND lote=%s", (pid, uid, str(v.get("lote") or "IMP")))
                cur.execute("INSERT INTO stock_lotes (org_id, producto_id, ubicacion_id, lote, cantidad, vencimiento) VALUES (%s,%s,%s,%s,%s,%s)",
                            (ctx.org_id, pid, uid, str(v.get("lote") or "IMP"), cant, fecha(v.get("vencimiento"))))
            n += 1
    return {"importadas": n, "mensaje": f"Stock actualizado en {n} productos × sucursal"}


def _productos(conn, ctx, lote, filas, cat, zona) -> dict:
    nuevos = actualizados = sin_mapear = 0
    categorias = {(_norm(c["nombre"]), c["padre_id"]): c["id"] for c in db.filas(conn, "SELECT id, nombre, padre_id FROM categorias")}
    hoy = datetime.now(zona).date()

    def categoria(nombre, padre=None):
        if not nombre:
            return padre
        clave = (_norm(nombre), padre)
        if clave not in categorias:
            categorias[clave] = db.fila(conn, "INSERT INTO categorias (org_id, nombre, padre_id, created_by) VALUES (%s,%s,%s,%s) RETURNING id",
                                        (ctx.org_id, str(nombre).strip(), padre, ctx.usuario_id))["id"]
        return categorias[clave]
    with conn.cursor() as cur:
        for _, v in filas:
            cat_id = categoria(v.get("categoria"))
            sub_id = categoria(v.get("subcategoria"), cat_id) if v.get("subcategoria") else cat_id
            ean = str(v.get("ean") or "").strip() or None
            maestro_id, estado = catalogo.vincular_maestro(conn, {"ean": ean})
            codigo = str(v["codigo"]).strip()
            unidad = "kg" if _norm(v.get("unidad") or "") in ("kg", "kilo", "kilos") else "unidad"
            existente = db.fila(conn, "SELECT id FROM productos WHERE codigo_interno=%s", (codigo,))
            if existente:
                pid = existente["id"]
                cur.execute("UPDATE productos SET nombre=%s, ean=coalesce(%s, ean), marca=coalesce(%s, marca), categoria_id=coalesce(%s, categoria_id), "
                            "maestro_id=coalesce(%s, maestro_id), estado_mapeo=CASE WHEN %s IS NOT NULL THEN 'mapeado' ELSE estado_mapeo END, updated_at=now() WHERE id=%s",
                            (str(v["nombre"]).strip(), ean, v.get("marca"), sub_id, maestro_id, maestro_id, pid))
                actualizados += 1
            else:
                pid = db.fila(conn, "INSERT INTO productos (org_id, maestro_id, codigo_interno, nombre, ean, marca, categoria_id, unidad, estado_mapeo, created_by) "
                                    "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id",
                              (ctx.org_id, maestro_id, codigo, str(v["nombre"]).strip(), ean, v.get("marca"), sub_id, unidad, estado, ctx.usuario_id))["id"]
                nuevos += 1
            sin_mapear += estado == "sin_mapear"
            if v.get("proveedor"):
                prov = cat["proveedores"].get(_norm(v["proveedor"]))
                if not prov:
                    prov = db.fila(conn, "INSERT INTO proveedores (org_id, razon_social, created_by) VALUES (%s,%s,%s) RETURNING id",
                                   (ctx.org_id, str(v["proveedor"]).strip(), ctx.usuario_id))["id"]
                    cat["proveedores"][_norm(v["proveedor"])] = prov
                cur.execute("INSERT INTO producto_proveedores (org_id, producto_id, proveedor_id, costo, unidades_por_bulto) VALUES (%s,%s,%s,%s,%s) "
                            "ON CONFLICT (producto_id, proveedor_id) DO UPDATE SET costo=coalesce(EXCLUDED.costo, producto_proveedores.costo), "
                            "unidades_por_bulto=EXCLUDED.unidades_por_bulto, updated_at=now()",
                            (ctx.org_id, pid, prov, numero(v.get("costo")), int(numero(v.get("bulto")) or 1)))
            precio = numero(v.get("precio"))
            if precio:
                vigente = db.fila(conn, "SELECT id, precio FROM precios WHERE producto_id=%s AND canal_id IS NULL AND ubicacion_id IS NULL AND hasta IS NULL", (pid,))
                if not vigente or vigente["precio"] != precio:
                    if vigente:
                        cur.execute("UPDATE precios SET hasta=%s WHERE id=%s", (hoy - timedelta(days=1), vigente["id"]))
                    cur.execute("INSERT INTO precios (org_id, producto_id, precio, desde, origen, created_by) VALUES (%s,%s,%s,%s,'importado',%s)",
                                (ctx.org_id, pid, precio, hoy, ctx.usuario_id))
    return {"importadas": nuevos + actualizados, "nuevos": nuevos, "actualizados": actualizados, "sin_mapear": sin_mapear,
            "mensaje": f"{nuevos} productos nuevos y {actualizados} actualizados" + (f"; {sin_mapear} con código de barras que el catálogo maestro no conoce (revisalos en Catálogo)" if sin_mapear else "")}


def _compras(conn, ctx, lote, filas, cat, zona) -> dict:
    from .api_documentos import LineaRecibida, registrar_recepcion
    grupos: dict = defaultdict(list)
    for _, v in filas:
        grupos[(_norm(v["proveedor"]), str(v["documento"]).strip(), cat["ubicaciones"][_norm(v["sucursal"])])].append(v)
    importadas = duplicadas = 0
    for (prov_n, documento, uid), vs in grupos.items():
        prov = cat["proveedores"].get(prov_n)
        if not prov:
            prov = db.fila(conn, "INSERT INTO proveedores (org_id, razon_social, created_by) VALUES (%s,%s,%s) RETURNING id",
                           (ctx.org_id, str(vs[0]["proveedor"]).strip(), ctx.usuario_id))["id"]
            cat["proveedores"][prov_n] = prov
        if db.fila(conn, "SELECT 1 FROM recepciones WHERE proveedor_id=%s AND documento=%s AND ubicacion_id=%s", (prov, documento, uid)):
            duplicadas += 1
            continue
        lineas = []
        for v in vs:
            pid, _ = _producto(conn, cat, v["producto"], "proveedor", prov)
            lineas.append(LineaRecibida(producto_id=pid, ubicacion_id=uid, cantidad=float(numero(v["cantidad"])), costo=str(numero(v["costo"])),
                                        lote=v.get("lote"), vencimiento=fecha(v.get("vencimiento"))))
        registrar_recepcion(conn, ctx, None, prov, uid, lineas, documento)
        importadas += 1
    return {"importadas": importadas, "duplicadas": duplicadas,
            "mensaje": f"{importadas} comprobantes de compra cargados" + (f"; {duplicadas} ya estaban" if duplicadas else "")}


def _precios(conn, ctx, lote, filas, cat, zona) -> dict:
    """Lista de precios del proveedor: se registra con vigencia, se empareja cada línea y se actualiza el costo vigente
    (eso dispara la remarcación y el aviso de aumento)."""
    por_prov: dict = defaultdict(list)
    for _, v in filas:
        por_prov[_norm(v["proveedor"])].append(v)
    emparejadas = sin_emparejar = 0
    pendientes = []
    for prov_n, vs in por_prov.items():
        prov = cat["proveedores"].get(prov_n)
        if not prov:
            prov = db.fila(conn, "INSERT INTO proveedores (org_id, razon_social, created_by) VALUES (%s,%s,%s) RETURNING id",
                           (ctx.org_id, str(vs[0]["proveedor"]).strip(), ctx.usuario_id))["id"]
        vigencia = fecha(vs[0].get("vigencia")) or datetime.now(zona).date()
        lista = db.fila(conn, "INSERT INTO listas_precios_proveedor (org_id, proveedor_id, vigencia_desde, origen, created_by) VALUES (%s,%s,%s,'archivo',%s) RETURNING id",
                        (ctx.org_id, prov, vigencia, ctx.usuario_id))["id"]
        with conn.cursor() as cur:
            for v in vs:
                codigo = str(v.get("producto") or "").strip() or None
                e = catalogo.emparejar(conn, ean=codigo if codigo and codigo.isdigit() and len(codigo) >= 8 else None, codigo=codigo,
                                       texto=v.get("descripcion"), origen="proveedor", origen_id=prov, productos=cat["productos"])
                costo = numero(v["costo"])
                anterior = None
                if e.producto_id:
                    a = db.fila(conn, "SELECT costo FROM producto_proveedores WHERE producto_id=%s AND proveedor_id=%s", (e.producto_id, prov))
                    anterior = a["costo"] if a else None
                    cur.execute("INSERT INTO producto_proveedores (org_id, producto_id, proveedor_id, costo, codigo_proveedor, principal) VALUES (%s,%s,%s,%s,%s,%s) "
                                "ON CONFLICT (producto_id, proveedor_id) DO UPDATE SET costo=EXCLUDED.costo, codigo_proveedor=coalesce(EXCLUDED.codigo_proveedor, producto_proveedores.codigo_proveedor), updated_at=now()",
                                (ctx.org_id, e.producto_id, prov, costo, codigo,
                                 not db.fila(conn, "SELECT 1 FROM producto_proveedores WHERE producto_id=%s AND principal", (e.producto_id,))))
                    emparejadas += 1
                else:
                    sin_emparejar += 1
                    pendientes.append({"codigo": codigo, "descripcion": v.get("descripcion"), "costo": costo, "candidatos": e.candidatos[:3]})
                cur.execute("INSERT INTO listas_precios_proveedor_lineas (org_id, lista_id, producto_id, codigo, descripcion, costo, costo_anterior) "
                            "VALUES (%s,%s,%s,%s,%s,%s,%s)", (ctx.org_id, lista, e.producto_id, codigo, v.get("descripcion"), costo, anterior))
    return {"importadas": emparejadas, "sin_emparejar": sin_emparejar, "pendientes": pendientes[:50],
            "mensaje": f"{emparejadas} costos actualizados" + (f"; {sin_emparejar} líneas sin emparejar (revisalas en Catálogo)" if sin_emparejar else "")}


# Modo distribuidor: clientes, pedidos y cuenta corriente (al final: ese módulo usa las funciones de este).
from . import importar_distribuidor  # noqa: E402

CAMPOS.update(importar_distribuidor.CAMPOS)
for _campo, _sinonimos in importar_distribuidor.SINONIMOS.items():
    SINONIMOS[_campo] = SINONIMOS.get(_campo, []) + [x for x in _sinonimos if x not in SINONIMOS.get(_campo, [])]
